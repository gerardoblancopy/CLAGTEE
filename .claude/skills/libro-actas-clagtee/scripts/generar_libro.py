#!/usr/bin/env python3
"""Genera el Libro de Resúmenes (Book of Abstracts) de CLAGTEE con el formato
del libro 2019.

Pasos:
  1. Lee los papers (export Excel del CMS, JSON de /api/papers o CSV).
  2. Se queda con los aceptados, los asigna a su track oficial, los numera y
     arma las entradas del índice de autores.
  3. Escribe libro.json y reporte.md (pendientes, advertencias, índice de
     autores para revisar) en la carpeta de salida.
  4. Compila plantilla/libro.typ con Typst y deja el PDF (y opcionalmente la
     versión en pliegos y vistas previas PNG).

Uso típico (desde la raíz del repo):
  python3 .claude/skills/libro-actas-clagtee/scripts/generar_libro.py \
      --papers ~/Downloads/clagtee2026_papers_todos-completo_20261105_1200.xlsx

Requisitos: Python 3.11+, `pip install typst openpyxl` (pypdf solo para
--pliegos).
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import subprocess
import sys
import tomllib
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parents[1]
REPO_ROOT = SKILL_DIR.parents[2]
DEFAULT_CONFIG = SKILL_DIR / "config" / "clagtee2026.toml"
TEMPLATE = SKILL_DIR / "plantilla" / "libro.typ"
FONTS_DIR = SKILL_DIR / "assets" / "fonts"
# Fuentes con licencia que no se puede redistribuir (Nirmala UI, Minion Pro,
# Modernica): se copian aquí y git las ignora.
FONTS_LOCAL = SKILL_DIR / "assets" / "fonts-locales"
DEFAULT_OUT = REPO_ROOT / "scratch" / "libro-actas"

ACCEPTED = {"accepted", "aceptado", "aceptada"}

# ---------------------------------------------------------------------------
# Utilidades de texto
# ---------------------------------------------------------------------------


def sin_tildes(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def clave_norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", sin_tildes(s).lower()).strip()


def espacios(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


PREFIJO_RESUMEN = re.compile(r"^\s*(abstract|resumen|resumo)\s*[—–:.\-]+\s*", re.I)


def limpiar_resumen(s: str) -> list[str]:
    s = str(s or "").replace("\r\n", "\n").strip()
    s = PREFIJO_RESUMEN.sub("", s)
    parrafos = [espacios(p) for p in re.split(r"\n\s*\n", s)]
    return [p for p in parrafos if p]


def limpiar_titulo(s: str) -> str:
    return espacios(s)


PALABRAS_IDIOMA = {
    "en": " the of and is this with for are that from to in on by an ",
    "es": " el la los las del que para una con se por como en es al este esta ",
    "pt": " o os do da dos das que para uma com não em ao no na pelo pela este esta ",
}


def detectar_idioma(texto: str) -> str:
    palabras = re.findall(r"[a-záéíóúâêôãõçñü]+", texto.lower())
    puntaje = {}
    for idioma, lista in PALABRAS_IDIOMA.items():
        vocab = set(lista.split())
        puntaje[idioma] = sum(1 for p in palabras if p in vocab)
    puntaje["pt"] += 3 * len(re.findall(r"ção|ções|ões|ã", texto.lower()))
    puntaje["es"] += 3 * len(re.findall(r"ción|ñ", texto.lower()))
    return max(puntaje, key=puntaje.get) if palabras else "en"


# ---------------------------------------------------------------------------
# Nombres de autores
# ---------------------------------------------------------------------------

TITULOS = re.compile(r"^(prof|profa|dr|dra|ing|eng|msc|m\.sc|mg|phd|ph\.d|lic|sr|sra)\.?\s+", re.I)
GRADOS = re.compile(r",\s*(ph\.?\s?d\.?|m\.?\s?sc\.?|msc|eng|ing|b\.?sc\.?|mba|jr\.?)\s*$", re.I)
PARTICULAS = {"de", "da", "do", "dos", "das", "del", "della", "di", "du", "van", "von", "la", "le", "y", "e"}
SUFIJOS = {"jr", "jr.", "junior", "júnior", "filho", "neto", "sobrinho", "fº", "f°", "segundo", "ii", "iii"}
PAISES_HISPANOS = {
    "cl", "ar", "co", "pe", "mx", "es", "uy", "py", "bo", "ec", "ve", "cr", "gt", "pa", "do", "cu", "hn", "sv", "ni",
}
NOMBRES_PAISES_HISPANOS = (
    "chile", "argentina", "colombia", "peru", "perú", "mexico", "méxico", "españa", "spain", "uruguay",
    "paraguay", "bolivia", "ecuador", "venezuela", "costa rica", "guatemala", "panama", "panamá",
)


def limpiar_nombre(nombre: str) -> str:
    n = espacios(nombre)
    n = TITULOS.sub("", n)
    n = GRADOS.sub("", n)
    if n.isupper() or n.islower():
        partes = []
        for p in n.split(" "):
            partes.append(p.lower() if p.lower() in PARTICULAS else p[:1].upper() + p[1:].lower())
        n = " ".join(partes)
    return n


def convencion(autor: dict) -> str:
    email = (autor.get("email") or "").lower()
    afil = (autor.get("afiliacion") or "").lower()
    tld = email.rsplit(".", 1)[-1] if "@" in email else ""
    if tld in {"br", "pt"} or "brasil" in afil or "brazil" in afil or "universidade" in afil:
        return "lusa"
    if tld in PAISES_HISPANOS or any(p in afil for p in NOMBRES_PAISES_HISPANOS) or "universidad" in afil:
        return "chilena" if (tld == "cl" or "chile" in afil) else "hispana"
    return "desconocida"


def dividir_nombre(autor: dict) -> tuple[str, str, str]:
    """Devuelve (apellidos, nombres, motivo) según la convención del autor."""
    nombre = autor["nombre"]
    if "," in nombre:
        ap, nom = [espacios(x) for x in nombre.split(",", 1)]
        return ap, nom, "formato 'Apellido, Nombre'"
    tokens = nombre.split(" ")
    if len(tokens) == 1:
        return tokens[0], "", "un solo nombre"
    conv = convencion(autor)
    motivo = "último apellido"
    es_inicial = re.compile(r"^\w\.?$")
    if tokens[-1].lower() in SUFIJOS and len(tokens) >= 3:
        # "Wilson R. D. Junior" -> "Junior, W"; "Ivo Silva Junior" -> "Silva Junior, I"
        inicio = len(tokens) - 1 if es_inicial.match(tokens[-2]) else len(tokens) - 2
        motivo = "apellido + sufijo"
    elif conv in {"hispana", "chilena"} and len(tokens) >= 4:
        inicio = len(tokens) - 2
        motivo = "dos apellidos (convención hispana)"
    elif conv == "chilena" and len(tokens) == 3:
        inicio = 1
        motivo = "AMBIGUO: 3 palabras, se asumió nombre + 2 apellidos"
    else:
        inicio = len(tokens) - 1
        if conv in {"hispana", "desconocida"} and len(tokens) == 3:
            motivo = "AMBIGUO: 3 palabras, se asumió 2 nombres + apellido"
    while inicio > 1 and tokens[inicio - 1].lower() in PARTICULAS:
        inicio -= 1
    return " ".join(tokens[inicio:]), " ".join(tokens[:inicio]), motivo


def entrada_indice(autor: dict, correcciones: dict) -> tuple[str, str]:
    """Devuelve ("Apellido, I", motivo) para el índice de autores."""
    if autor["nombre"] in correcciones:
        return correcciones[autor["nombre"]], "corrección manual"
    apellido, nombres, motivo = dividir_nombre(autor)
    return (f"{apellido}, {nombres[:1].upper()}" if nombres else apellido), motivo


def unir_autores(nombres: list[str], conector: str) -> str:
    if len(nombres) <= 1:
        return "".join(nombres)
    return ", ".join(nombres[:-1]) + f" {conector} " + nombres[-1]


# ---------------------------------------------------------------------------
# Lectura de papers
# ---------------------------------------------------------------------------

AUTOR_DETALLE = re.compile(
    r"^(?P<nombre>.*?)(?:\s*<(?P<email>[^>]*)>)?(?:\s*\[(?P<afil>[^\]]*)\])?(?:\s*\(Correspondiente\))?\s*$"
)

ALIAS = {
    "id": ["ID Paper", "id", "ID", "paper_id"],
    "titulo": ["Título", "Titulo", "title", "titulo"],
    "titulo_en": ["Título en Inglés (IEEE)", "titleEn", "title_en"],
    "resumen": ["Abstract (Envío Original)", "Abstract", "Resumen", "abstract", "resumen"],
    "resumen_en": ["Abstract en Inglés (IEEE)", "abstractEn", "abstract_en"],
    "track": ["Área Temática (Track)", "track", "Track"],
    "estado": ["Clasificación / Estado", "status", "estado", "Estado"],
    "autores_detalle": ["Detalle Completo Autores"],
    "autores": ["authors", "autores", "Autores"],
    "inscripcion": ["Tiene Inscripción en CMS"],
    "subtema": ["subtema", "subtopic", "Subtema"],
    "idioma": ["idioma", "lang", "language"],
}


def campo(fila: dict, nombre: str):
    for k in ALIAS[nombre]:
        if k in fila and fila[k] not in (None, ""):
            return fila[k]
    return None


def autores_de(fila: dict) -> list[dict]:
    crudo = campo(fila, "autores")
    if isinstance(crudo, list):
        out = []
        for a in crudo:
            if isinstance(a, str):
                out.append({"nombre": a})
            elif isinstance(a, dict):
                out.append({"nombre": a.get("name") or a.get("nombre") or "", "email": a.get("email", ""),
                            "afiliacion": a.get("affiliation") or a.get("afiliacion") or ""})
        return [a for a in out if espacios(a["nombre"])]
    detalle = campo(fila, "autores_detalle")
    if detalle:
        out = []
        for parte in str(detalle).split(" | "):
            m = AUTOR_DETALLE.match(parte.strip())
            if m and espacios(m.group("nombre")):
                out.append({"nombre": m.group("nombre"), "email": m.group("email") or "", "afiliacion": m.group("afil") or ""})
        return out
    if isinstance(crudo, str):
        return [{"nombre": n} for n in re.split(r"\s*;\s*", crudo) if espacios(n)]
    principal = fila.get("Autor Correspondiente (Quien Sometió)")
    coautores = fila.get("Nombres de Co-autores") or ""
    nombres = ([principal] if principal else []) + [n for n in str(coautores).split(";") if espacios(n)]
    return [{"nombre": n} for n in nombres]


def leer_xlsx(ruta: Path) -> list[dict]:
    try:
        import openpyxl
    except ImportError:
        sys.exit("Falta openpyxl: pip install openpyxl")
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    hojas = wb.sheetnames
    preferida = next((h for h in hojas if h.lower().startswith("aceptado")), None)
    preferida = preferida or next((h for h in hojas if h.lower().startswith(("todos", "papers"))), hojas[0])
    filas = list(wb[preferida].iter_rows(values_only=True))
    if not filas:
        return []
    cab = [str(c or "").strip() for c in filas[0]]
    return [dict(zip(cab, f)) for f in filas[1:] if any(v not in (None, "") for v in f)]


def leer_papers(ruta: Path) -> list[dict]:
    suf = ruta.suffix.lower()
    if suf in {".xlsx", ".xlsm"}:
        return leer_xlsx(ruta)
    if suf == ".csv":
        with ruta.open(encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))
    if suf == ".json":
        datos = json.loads(ruta.read_text(encoding="utf-8"))
        return datos["papers"] if isinstance(datos, dict) and "papers" in datos else datos
    sys.exit(f"Formato no soportado: {ruta.suffix} (use .xlsx, .csv o .json)")


PAISES_TLD = {
    "br": "Brasil", "cl": "Chile", "ar": "Argentina", "pe": "Perú", "co": "Colombia", "mx": "México",
    "uy": "Uruguay", "py": "Paraguay", "bo": "Bolivia", "ec": "Ecuador", "ve": "Venezuela", "es": "España",
    "pt": "Portugal", "uk": "Reino Unido", "de": "Alemania", "fr": "Francia", "it": "Italia", "no": "Noruega",
    "cu": "Cuba", "hn": "Honduras", "cr": "Costa Rica", "in": "India", "hk": "Hong Kong", "ca": "Canadá",
}
PAISES_TEXTO = {
    "brasil": "Brasil", "brazil": "Brasil", "chile": "Chile", "argentina": "Argentina", "peru": "Perú",
    "perú": "Perú", "colombia": "Colombia", "mexico": "México", "méxico": "México", "uruguay": "Uruguay",
    "paraguay": "Paraguay", "bolivia": "Bolivia", "ecuador": "Ecuador", "venezuela": "Venezuela",
    "españa": "España", "spain": "España", "portugal": "Portugal", "usa": "Estados Unidos",
    "united states": "Estados Unidos", "estados unidos": "Estados Unidos", "united kingdom": "Reino Unido",
    "germany": "Alemania", "alemania": "Alemania", "france": "Francia", "italy": "Italia", "norway": "Noruega",
    "canada": "Canadá", "india": "India", "cuba": "Cuba", "honduras": "Honduras",
}
# Siglas de instituciones frecuentes en el congreso, por si la afiliación
# no nombra el país y el email es genérico (gmail, outlook...).
INSTITUCIONES = {
    "pucv": "Chile", "usach": "Chile", "utfsm": "Chile", "unab": "Chile", "uchile": "Chile", "puc": "Chile",
    "udec": "Chile", "ufro": "Chile", "unesp": "Brasil", "usp": "Brasil", "unicamp": "Brasil", "utfpr": "Brasil",
    "ufjf": "Brasil", "ufpr": "Brasil", "ufpa": "Brasil", "ufam": "Brasil", "ufg": "Brasil", "unifei": "Brasil",
    "ufrj": "Brasil", "cepel": "Brasil", "unmdp": "Argentina", "utn": "Argentina", "unsj": "Argentina",
    "conicet": "Argentina", "uba": "Argentina", "unlp": "Argentina",
}


def pais_de(email: str, afiliacion: str) -> str:
    afil = (afiliacion or "").lower()
    for clave, pais in PAISES_TEXTO.items():
        if re.search(rf"\b{re.escape(clave)}\b", afil):
            return pais
    for sigla, pais in INSTITUCIONES.items():
        if re.search(rf"\b{sigla}\b", afil):
            return pais
    dominio = (email or "").lower().rsplit("@", 1)[-1]
    tld = dominio.rsplit(".", 1)[-1] if "." in dominio else ""
    if tld in PAISES_TLD:
        return PAISES_TLD[tld]
    for sigla, pais in INSTITUCIONES.items():
        if sigla in dominio.split("."):
            return pais
    return ""


def leer_revisores(ruta: Path) -> list[dict] | None:
    """Hoja "Revisores (N)" del Excel del CMS, o un CSV/XLSX con columnas
    Nombre, Email, Afiliación. None si el archivo no trae revisores."""
    filas: list[dict] = []
    suf = ruta.suffix.lower()
    if suf in {".xlsx", ".xlsm"}:
        import openpyxl
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        hoja = next((h for h in wb.sheetnames if h.lower().startswith(("revisores", "reviewers"))), None)
        if hoja is None:
            return None
        datos = list(wb[hoja].iter_rows(values_only=True))
        cab = [str(c or "").strip() for c in datos[0]] if datos else []
        filas = [dict(zip(cab, f)) for f in datos[1:] if any(v not in (None, "") for v in f)]
    elif suf == ".csv":
        with ruta.open(encoding="utf-8-sig", newline="") as f:
            filas = list(csv.DictReader(f))
    elif suf == ".json":
        # JSON de GET /api/auth/users?role=reviewer: {"users": [...]}
        datos = json.loads(ruta.read_text(encoding="utf-8"))
        if not (isinstance(datos, dict) and isinstance(datos.get("users"), list)):
            return None
        filas = [u for u in datos["users"] if isinstance(u, dict) and u.get("role", "reviewer") == "reviewer"]
    out = []
    for f in filas:
        nombre = espacios(f.get("Nombre") or f.get("name") or f.get("nombre") or "")
        if not nombre:
            continue
        evals = f.get("Evaluaciones Enviadas")
        out.append({
            "nombre": limpiar_nombre(nombre),
            "email": espacios(f.get("Email") or f.get("email") or ""),
            "afiliacion": espacios(f.get("Afiliación") or f.get("affiliation") or f.get("afiliacion") or ""),
            "evaluaciones": int(evals) if str(evals or "").strip().isdigit() else None,
        })
    return out


def comite_desde_revisores(revisores: list[dict], cfg: dict, solo_con_evaluaciones: bool) -> tuple[list, list[str]]:
    correcciones = cfg.get("comite_programa", {}).get("correcciones", {})
    avisos = []
    miembros = {}
    for r in revisores:
        if solo_con_evaluaciones and r["evaluaciones"] is not None and r["evaluaciones"] == 0:
            continue
        if r["nombre"] in correcciones:
            entrada, motivo = correcciones[r["nombre"]], "corrección manual"
        else:
            apellido, nombres, motivo = dividir_nombre(r)
            entrada = f"{apellido}, {nombres}" if nombres else apellido
        pais = pais_de(r["email"], r["afiliacion"])
        if not pais:
            avisos.append(f"Comité de Programa: sin país para '{r['nombre']}' ({r['afiliacion'] or r['email']}).")
        if motivo.startswith("AMBIGUO"):
            avisos.append(f"Comité de Programa: '{r['nombre']}' -> '{entrada}' ({motivo}).")
        miembros[clave_norm(entrada)] = [entrada, pais]
    lista = sorted(miembros.values(), key=lambda m: clave_norm(m[0]))
    return lista, avisos


# ---------------------------------------------------------------------------
# Armado del libro
# ---------------------------------------------------------------------------


def numero_id(pid: str) -> tuple:
    m = re.search(r"(\d+)\s*$", pid)
    return (int(m.group(1)) if m else 10**9, pid)


def etiqueta(prefijo: str, s: str, usadas: set) -> str:
    base = prefijo + re.sub(r"[^A-Za-z0-9_.:-]", "_", s)
    e, k = base, 2
    while e in usadas:
        e, k = f"{base}-{k}", k + 1
    usadas.add(e)
    return e


def buscar_todos(obj, ruta=""):
    if isinstance(obj, dict):
        for k, v in obj.items():
            yield from buscar_todos(v, f"{ruta}.{k}" if ruta else k)
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from buscar_todos(v, f"{ruta}[{i}]")
    elif isinstance(obj, str) and obj.strip().startswith("TODO"):
        yield ruta, obj.strip()


def construir(args, cfg: dict) -> tuple[dict, list[str], list[tuple]]:
    filas = leer_papers(Path(args.papers).expanduser())
    avisos: list[str] = []
    excluir = set(cfg.get("seleccion", {}).get("excluir", []))
    tracks_cfg = cfg["tracks"]
    mapa_track = {}
    for tr in tracks_cfg:
        for k in ("es", "en", "pt"):
            mapa_track[clave_norm(tr[k])] = tr
        mapa_track[clave_norm(f"track {tr['numero']}")] = tr
        mapa_track[str(tr["numero"])] = tr
    conector = cfg["libro"].get("conector_autores", "and")
    formato_id = cfg["libro"].get("formato_id", "(ID {id})")

    por_track: dict[int, list[dict]] = defaultdict(list)
    vistos_id = set()
    for fila in filas:
        pid = espacios(campo(fila, "id") or "")
        estado = clave_norm(str(campo(fila, "estado") or ""))
        if not pid:
            continue
        if estado and estado not in ACCEPTED and not args.incluir_todos:
            continue
        if pid in excluir:
            avisos.append(f"{pid}: excluido por config ([seleccion].excluir).")
            continue
        if pid in vistos_id:
            avisos.append(f"{pid}: ID repetido en el archivo; se usa la primera aparición.")
            continue
        vistos_id.add(pid)
        if args.requiere_inscripcion:
            ins = campo(fila, "inscripcion")
            if ins is not None and clave_norm(str(ins)) not in {"si", "yes", "true", "1"}:
                avisos.append(f"{pid}: sin inscripción en el CMS; omitido (--requiere-inscripcion).")
                continue

        titulo_en, resumen_en = campo(fila, "titulo_en"), campo(fila, "resumen_en")
        titulo_or, resumen_or = campo(fila, "titulo"), campo(fila, "resumen")
        if args.idioma == "en":
            titulo = titulo_en or titulo_or
            resumen = resumen_en or resumen_or
            if not titulo_en:
                avisos.append(f"{pid}: sin título en inglés confirmado; se usa el título original.")
            if not resumen_en and resumen_or:
                avisos.append(f"{pid}: sin abstract en inglés confirmado; se usa el abstract original.")
        else:
            titulo = titulo_or or titulo_en
            resumen = resumen_or or resumen_en
        titulo = limpiar_titulo(titulo or "")
        if titulo.isupper() and len(titulo) > 20:
            avisos.append(f"{pid}: título todo en MAYÚSCULAS; conviene corregirlo en el CMS.")
        parrafos = limpiar_resumen(resumen or "")
        if not titulo:
            titulo = "TODO: título pendiente"
            avisos.append(f"{pid}: SIN TÍTULO.")
        if not parrafos:
            parrafos = ["TODO: abstract pendiente."]
            avisos.append(f"{pid}: SIN ABSTRACT; se imprime un marcador en rojo.")

        track_crudo = espacios(campo(fila, "track") or "")
        tr = mapa_track.get(clave_norm(track_crudo)) or mapa_track.get(track_crudo)
        if not tr:
            tr = {"numero": 99, "en": "Other Contributions"}
            avisos.append(f"{pid}: track no reconocido '{track_crudo}'; va a 'Other Contributions'.")

        autores = []
        for a in autores_de(fila):
            a["nombre"] = limpiar_nombre(a["nombre"])
            autores.append(a)
        if not autores:
            avisos.append(f"{pid}: SIN AUTORES.")

        idioma = campo(fila, "idioma") or detectar_idioma(titulo + " " + " ".join(parrafos))
        por_track[tr["numero"]].append({
            "id": pid,
            "titulo": titulo,
            "resumen": parrafos,
            "autores": autores,
            "idioma": idioma,
            "subtema": espacios(campo(fila, "subtema") or "") or None,
            "track": tr,
        })

    if not por_track:
        sys.exit("No hay papers aceptados en el archivo (revise la hoja o use --incluir-todos).")
    if not any(campo(f, "resumen") or campo(f, "resumen_en") for f in filas):
        avisos.insert(0, "EL ARCHIVO NO TRAE ABSTRACTS (ni originales ni en inglés). Exporte de nuevo "
                         "desde el CMS: el Excel incluye 'Abstract (Envío Original)' y 'Abstract en Inglés (IEEE)'.")

    usadas: set = set()
    tracks_out = []
    secuencial = cfg.get("seleccion", {}).get("numeracion_tracks", "oficial") == "secuencial"
    correcciones = cfg.get("indice_autores", {}).get("correcciones", {})
    indice: dict[str, list] = defaultdict(list)
    revision: list[tuple] = []
    for k, num in enumerate(sorted(por_track), start=1):
        papers = sorted(por_track[num], key=lambda p: numero_id(p["id"]))
        # Subtemas: en el orden de su primera aparición, como en 2019.
        orden_sub = []
        for p in papers:
            if p["subtema"] and p["subtema"] not in orden_sub:
                orden_sub.append(p["subtema"])
        if orden_sub:
            papers.sort(key=lambda p: (orden_sub.index(p["subtema"]) if p["subtema"] else -1, numero_id(p["id"])))
        n_track = k if secuencial else num
        items = []
        sub_actual = None
        for j, p in enumerate(papers, start=1):
            if p["subtema"] and p["subtema"] != sub_actual:
                sub_actual = p["subtema"]
                items.append({"tipo": "subtema", "titulo": sub_actual, "etiqueta": etiqueta("s-", f"{n_track}-{sub_actual}", usadas)})
            sufijo = formato_id.format(id=p["id"])
            base = p["titulo"].rstrip()
            base = base if base[-1:] in ".?!" else base + "."
            et = etiqueta("p-", p["id"], usadas)
            nombres = [a["nombre"] for a in p["autores"]]
            items.append({
                "tipo": "paper",
                "id": p["id"],
                "numero": f"{n_track}.{j}",
                "titulo": base,
                "sufijo": sufijo,
                "autores": unir_autores(nombres, conector),
                "resumen": p["resumen"],
                "idioma": p["idioma"],
                "etiqueta": et,
            })
            for a in p["autores"]:
                clave, motivo = entrada_indice(a, correcciones)
                indice[clave].append(et)
                revision.append((a["nombre"], clave, motivo, p["id"], a.get("email", ""), a.get("afiliacion", "")))
        tracks_out.append({
            "numero": n_track,
            "titulo": p["track"]["en"],
            "etiqueta": etiqueta("t-", str(n_track), usadas),
            "items": items,
        })

    # Fusiona variantes con y sin tildes ("Aguero, C" / "Agüero, C").
    grupos: dict[str, list[str]] = defaultdict(list)
    for clave in indice:
        grupos[clave_norm(clave)].append(clave)
    entradas = []
    for norm, claves in grupos.items():
        conteo = Counter({c: len(indice[c]) for c in claves})
        mostrar = max(claves, key=lambda c: (conteo[c], c != sin_tildes(c)))
        etqs = []
        for c in claves:
            etqs.extend(indice[c])
        if len(claves) > 1:
            avisos.append(f"Índice de autores: se fusionaron {claves} como '{mostrar}'.")
        entradas.append({"clave": mostrar, "orden": norm, "etiquetas": list(dict.fromkeys(etqs))})
    entradas.sort(key=lambda e: e["orden"])
    letras = []
    for e in entradas:
        letra = sin_tildes(e["clave"][:1]).upper() or "#"
        if not letras or letras[-1]["letra"] != letra:
            letras.append({"letra": letra, "entradas": []})
        letras[-1]["entradas"].append({"clave": e["clave"], "etiquetas": e["etiquetas"]})

    total = sum(len([i for i in t["items"] if i["tipo"] == "paper"]) for t in tracks_out)
    datos = {
        "tracks": tracks_out,
        "autores": letras,
        "estadisticas": {"papers": total, "tracks": len(tracks_out), "entradas_indice_autores": len(entradas)},
    }
    return datos, avisos, revision


def escribir_reporte(salida: Path, cfg: dict, datos: dict, avisos: list[str], revision: list[tuple]) -> Path:
    pendientes = [p for p in buscar_todos(cfg)
                  if not (datos.get("comite_programa") and p[0].startswith("comite_programa"))]
    lineas = ["# Reporte del Libro de Resúmenes", ""]
    e = datos["estadisticas"]
    lineas += [f"- Papers: **{e['papers']}** en **{e['tracks']}** tracks", f"- Entradas del índice de autores: {e['entradas_indice_autores']}", ""]
    lineas += ["## Por track", ""]
    for t in datos["tracks"]:
        n = len([i for i in t["items"] if i["tipo"] == "paper"])
        lineas.append(f"- {t['numero']}. {t['titulo']}: {n}")
    lineas += ["", f"## Pendientes en la configuración ({len(pendientes)})", ""]
    lineas += [f"- `{r}`: {v}" for r, v in pendientes] or ["- Ninguno."]
    lineas += ["", f"## Advertencias de datos ({len(avisos)})", ""]
    lineas += [f"- {a}" for a in avisos] or ["- Ninguna."]
    ambiguos = [r for r in revision if r[2].startswith("AMBIGUO")]
    lineas += ["", f"## Índice de autores: entradas ambiguas ({len(ambiguos)})", "",
               "Revisar y, si corresponde, agregar en `[indice_autores.correcciones]` del .toml.", "",
               "| Nombre en el CMS | Entrada | Motivo | Paper |", "|---|---|---|---|"]
    lineas += [f"| {n} | {c} | {m} | {p} |" for n, c, m, p, *_ in ambiguos]
    ruta = salida / "reporte.md"
    ruta.write_text("\n".join(lineas) + "\n", encoding="utf-8")
    with (salida / "indice-autores.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["nombre_cms", "entrada_indice", "motivo", "paper", "email", "afiliacion"])
        w.writerows(sorted(set(revision), key=lambda r: (clave_norm(r[1]), r[3])))
    return ruta


# ---------------------------------------------------------------------------
# Compilación
# ---------------------------------------------------------------------------


def rel(p: Path) -> str:
    return "/" + p.resolve().relative_to(REPO_ROOT).as_posix()


# Variantes que el libro necesita de cada familia original (peso, estilo).
# Sin la negrita, Typst dibujaría en regular los títulos y autores.
VARIANTES = {
    "sans": [(400, "normal"), (700, "normal")],
    "serif": [(400, "normal")],
    "display": [(700, "normal")],
}


def variantes_instaladas(carpetas: list[Path]) -> dict[str, set] | None:
    """{familia en minúsculas: {(peso, estilo), ...}} visibles para Typst."""
    try:
        import typst
        fuentes = typst.Fonts(include_system_fonts=True, font_paths=[str(c) for c in carpetas]).fonts()
    except Exception:
        return None
    out: dict[str, set] = defaultdict(set)
    for f in fuentes:
        out[f.family.lower()].add((f.weight, f.style))
    return out


def elegir_fuentes(cfg: dict, carpetas: list[Path]) -> dict[str, str]:
    """Familias y factor de cuerpo por tipo. La fuente original del 2019 se
    usa solo si están todas sus variantes; si no, la sustituta libre con el
    factor que iguala el ancho del texto."""
    fuentes = cfg.get("fuentes", {})
    sustituta = fuentes.get("escala_sustituta", {})
    instaladas = variantes_instaladas(carpetas)
    out = {}
    for tipo, necesarias in VARIANTES.items():
        lista = fuentes.get(tipo) or []
        original = lista[0] if lista else ""
        tiene = instaladas.get(original.lower(), set()) if instaladas is not None else set()
        faltan = [v for v in necesarias if v not in tiene]
        if lista and not faltan:
            familias, escala = lista, "1"
            print(f"  Fuente {tipo}: {original} (original)")
        else:
            familias, escala = lista[1:] or lista, str(sustituta.get(tipo, 1))
            if tiene:
                nombres = ", ".join("negrita" if peso >= 700 else "regular" for peso, _ in faltan)
                print(f"  AVISO: {original} está instalada pero le falta: {nombres}. Se usa {familias[0]}.")
            print(f"  Fuente {tipo}: {familias[0]} (sustituta libre, escala {escala})")
        out[f"fuentes-{tipo}"] = "|".join(familias)
        out[f"escala-{tipo}"] = escala
        # Sin cursiva en la familia elegida, la plantilla inclina el texto.
        cursiva = instaladas is not None and any(e == "italic" for _, e in instaladas.get(familias[0].lower(), set()))
        out[f"cursiva-{tipo}"] = "1" if cursiva or instaladas is None else "0"
    return out


def compilar(config: Path, cfg: dict, datos_json: Path, pdf: Path, carpetas: list[Path]) -> None:
    entradas = {"config": rel(config), "data": rel(datos_json), **elegir_fuentes(cfg, carpetas)}
    try:
        import typst
    except ImportError:
        typst = None
    if typst is not None:
        pdf_bytes, avisos = typst.compile_with_warnings(
            str(TEMPLATE), root=str(REPO_ROOT), font_paths=[str(c) for c in carpetas], sys_inputs=entradas)
        for a in avisos:
            msg = getattr(a, "message", str(a))
            if "unknown font family" not in msg:
                print(f"  typst: {msg}")
        pdf.write_bytes(pdf_bytes)
        return
    cli = shutil.which("typst")
    if not cli:
        sys.exit("Falta Typst: pip install typst (o instale la CLI de typst).")
    cmd = [cli, "compile", "--root", str(REPO_ROOT)]
    for c in carpetas:
        cmd += ["--font-path", str(c)]
    for k, v in entradas.items():
        cmd += ["--input", f"{k}={v}"]
    subprocess.run(cmd + [str(TEMPLATE), str(pdf)], check=True)


def pliegos(pdf: Path, destino: Path) -> None:
    """Versión para pantalla como el PDF 2019: portada sola, páginas
    enfrentadas (par | impar) y contraportada sola."""
    from pypdf import PdfReader, PdfWriter, Transformation
    r = PdfReader(str(pdf))
    w = PdfWriter()
    n = len(r.pages)
    ancho = float(r.pages[0].mediabox.width)
    alto = float(r.pages[0].mediabox.height)
    w.add_page(r.pages[0])
    for i in range(1, n - 1, 2):
        hoja = w.add_blank_page(width=2 * ancho, height=alto)
        hoja.merge_page(r.pages[i])
        if i + 1 < n - 1:
            hoja.merge_transformed_page(r.pages[i + 1], Transformation().translate(tx=ancho, ty=0))
    w.add_page(r.pages[n - 1])
    with destino.open("wb") as f:
        w.write(f)


def vistas_previas(pdf: Path, paginas: str, carpeta: Path) -> None:
    carpeta.mkdir(parents=True, exist_ok=True)
    if not shutil.which("pdftoppm"):
        print("  (pdftoppm no está instalado; se omiten las vistas previas)")
        return
    try:
        from pypdf import PdfReader
        total = len(PdfReader(str(pdf)).pages)
    except ImportError:
        total = 10**6
    for p in [x.strip() for x in paginas.split(",") if x.strip()]:
        a, _, b = p.partition("-")
        a, b = int(a), min(int(b or a), total)
        if a > b:
            print(f"  (página {p} fuera del libro de {total} páginas; se omite)")
            continue
        subprocess.run(["pdftoppm", "-r", "80", "-png", "-f", str(a), "-l", str(b), str(pdf), str(carpeta / "pag")], check=True)
    print(f"  Vistas previas en {carpeta}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--papers", required=True, help="Export del CMS (.xlsx), JSON de /api/papers o CSV")
    ap.add_argument("--config", default=str(DEFAULT_CONFIG), help="Archivo .toml del libro")
    ap.add_argument("--salida", default=str(DEFAULT_OUT), help="Carpeta de salida (dentro del repo)")
    ap.add_argument("--nombre", default="CLAGTEE2026_Book_of_Abstracts", help="Nombre base del PDF")
    ap.add_argument("--idioma", choices=["en", "original"], default="en",
                    help="en: título/abstract en inglés confirmados para IEEE (por defecto); original: los del envío")
    ap.add_argument("--requiere-inscripcion", action="store_true", help="Omite papers sin inscripción en el CMS")
    ap.add_argument("--incluir-todos", action="store_true", help="No filtra por estado 'aceptado' (pruebas)")
    ap.add_argument("--solo-datos", action="store_true", help="Genera libro.json y reporte.md sin compilar")
    ap.add_argument("--pliegos", action="store_true", help="Genera además la versión en pliegos (como el PDF 2019)")
    ap.add_argument("--vista-previa", default="", help="Páginas a exportar como PNG, p. ej. '1,3-5,27'")
    ap.add_argument("--revisores", default="",
                    help="Archivo con los revisores (por defecto, la hoja 'Revisores' del mismo Excel)")
    ap.add_argument("--revisores-con-evaluaciones", action="store_true",
                    help="Comité de Programa solo con revisores que enviaron al menos una evaluación")
    ap.add_argument("--fuentes", action="append", default=[],
                    help="Carpeta extra con fuentes (.ttf/.otf/.ttc); se puede repetir")
    args = ap.parse_args()

    config = Path(args.config).resolve()
    cfg = tomllib.loads(config.read_text(encoding="utf-8"))
    salida = Path(args.salida).resolve()
    try:
        salida.relative_to(REPO_ROOT)
    except ValueError:
        sys.exit(f"La carpeta de salida debe estar dentro del repositorio ({REPO_ROOT}): Typst solo lee archivos bajo su raíz.")
    salida.mkdir(parents=True, exist_ok=True)

    datos, avisos, revision = construir(args, cfg)
    fuente_rev = Path(args.revisores or args.papers).expanduser()
    revisores = leer_revisores(fuente_rev)
    if revisores:
        datos["comite_programa"], avisos_rev = comite_desde_revisores(revisores, cfg, args.revisores_con_evaluaciones)
        avisos.extend(avisos_rev)
        print(f"Comité de Programa: {len(datos['comite_programa'])} revisores del CMS ({fuente_rev.name}).")
    elif args.revisores:
        sys.exit(f"{fuente_rev} no trae una hoja o lista de revisores.")
    else:
        avisos.insert(0, "El archivo no trae la hoja 'Revisores': el Comité de Programa sale de la config "
                         "([comite_programa].miembros). Exporte de nuevo desde el CMS para usar los revisores.")
    datos_json = salida / "libro.json"
    datos_json.write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")
    reporte = escribir_reporte(salida, cfg, datos, avisos, revision)
    e = datos["estadisticas"]
    print(f"Papers: {e['papers']} en {e['tracks']} tracks; {len(avisos)} advertencias.")
    print(f"Reporte: {reporte}")
    if args.solo_datos:
        return

    pdf = salida / f"{args.nombre}.pdf"
    print("Compilando con Typst...")
    carpetas = [Path(c).expanduser().resolve() for c in args.fuentes] + [FONTS_LOCAL, FONTS_DIR]
    compilar(config, cfg, datos_json, pdf, [c for c in carpetas if c.is_dir()])
    print(f"PDF: {pdf}")
    if args.pliegos:
        destino = salida / f"{args.nombre}_pliegos.pdf"
        pliegos(pdf, destino)
        print(f"PDF en pliegos: {destino}")
    if args.vista_previa:
        vistas_previas(pdf, args.vista_previa, salida / "vista-previa")


if __name__ == "__main__":
    main()
